"""Write 9 Sep weekly outline, script, and Claims Desk card."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt, RGBColor

DESK = Path(r"C:\Users\BobbyWhittaker\OneDrive - Cassius Capital\Desktop")
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
MUTED = RGBColor(0x44, 0x44, 0x44)
INK = RGBColor(0x1A, 0x1A, 0x1A)


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = MUTED
    pf = style.paragraph_format
    pf.space_after = Pt(8)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = NAVY
    run.font.name = "Calibri"


def add_h(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = NAVY
    run.font.name = "Calibri"


def add_body(doc: Document, text: str, *, ink: bool = False) -> None:
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.font.color.rgb = INK if ink else MUTED


def add_quote(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = INK


def write_outline() -> Path:
    doc = Document()
    style_doc(doc)
    add_title(doc, "Roz weekly outline — 9 September 2026")
    add_body(
        doc,
        "Coverage: noon 31 August talk through 8 September. Rank IC book "
        "is unchanged: asof, generated_at=2026-08-17T17:28:40+00:00. v1 "
        "Claims Desk split asof-path-id-17aug-book-v1, same stamp. NVIDIA "
        "gold stamp 2026-08-27T18:02:00+00:00. Ops stamp "
        "2026-08-31T14:18:00+00:00. Healthcare desk stamp "
        "2026-08-31T14:45:00+00:00. Transparency sidecar "
        "2026-09-01T15:00:00+00:00. Quant/expire sidecar "
        "2026-09-01T16:00:00+00:00. Scorecard sidecar "
        "2026-09-09T01:27:26Z, 1,712 rows. Not a promotion. No holdout. "
        "No healthcare Rank IC. No production_v1.",
    )
    add_h(doc, "Through-line (~20 minutes)")
    add_body(
        doc,
        "Last talk, noon 31 August: a tree is not a row. Three books on "
        "Roz. NVIDIA gold deliver 6 / 7, cue recall 1.0. Tech and "
        "healthcare were first-seed books, leftover dump ~1,250, depth "
        "cohort paused. The room’s bar was B — tighten the pull — then C. "
        "This week B ran. After the room, honest desk and the leftover "
        "cull landed, then two typed depth closes. Then the pull became "
        "a machine: 310 proposed seeds, a 351-row review workbook, "
        "Desk Autopilot on every new call. Ops / healthcare grew from "
        "44 skeleton trees to hundreds of walked trees. The live object "
        "is Claims Desk + Post-Call Brief + a scorecard that plots every "
        "tracked call. The case study is DDOG and LITE — process, "
        "outcomes, meaning. CRWV stays on the desk as the younger-name "
        "precedent. Delivery on a first-seed book is still an em dash. "
        "Rank IC book did not move.",
    )
    add_h(doc, "1. What the room already heard (90 seconds)")
    add_body(
        doc,
        "Do not re-teach term structure. One sentence: same 17 August "
        "asof book, 0_56 is still not the typical quarter, not a "
        "promotion. Then the 31 August object, because this week only "
        "makes sense on top of it.",
    )
    add_body(
        doc,
        "v2 is a tree. One seed cite. Later quarters are edges. Walk "
        "may add restated or silent. Walk does not invent delivered, "
        "hit, or missed. Three books: nvda_gold_v2, desk_ops_v2, "
        "desk_hc_v2. Gold: 33 trees, 24 promises, 9 goals, window "
        "FY2022-Q2–FY2027-Q1, deliver 6 / 7 (H20 miss FY2026-Q3), "
        "hit 1 / 1, cue recall 18 / 18. v1 Flex kept + delivered still "
        "passes. Tech ops and healthcare were first-seed books. Deliver "
        "was an em dash, not 0%. Leftover cover was the hole.",
    )
    add_h(doc, "2. The afternoon after the talk — A rode, C started")
    add_body(
        doc,
        "Honest desk (option A) closed the same afternoon. Sector now "
        "drives Book: healthcare → desk_hc_v2, tech → desk_ops_v2, "
        "gold on All Companies or {NVDA}. Zero-scoreable rates caption "
        "as “No scored X yet. Em dash is not a 0% keep rate.” Trailing "
        "credibility stays NVIDIA-only.",
    )
    add_body(
        doc,
        "Tighten-the-pull cull, same afternoon: leftover openings "
        "1,250 → 1,145 (ops 610 → 556, healthcare 640 → 589). Shorter "
        "worklist, not a short one. Gold cue recall stayed 1.0. No "
        "auto-insert that day.",
        ink=True,
    )
    add_body(
        doc,
        "Depth cohort resumed after being paused in the talk. Typed "
        "closes from novelty_view only. ISRG da Vinci X — delivered "
        "FY2017-Q2 (11 of 166 systems were X; first clinical use in "
        "Germany; inside the FY2017-Q4 clock). Healthcare scored "
        "deliver 1 / 1 on that tree. CRM $20B next goal — restated "
        "FY2018-Q2, hit FY2021-Q4 (“over $20 billion in revenue”). "
        "The FY2021-Q1 guide is not the hit. Ops scored hit 1 / 1. "
        "LLY December 2016 dividend still unresolved. MSFT BUILD "
        "briefing still open. Do not invent those two.",
    )
    add_body(
        doc,
        "Same-object want→will is one tree. Display “became a promise.” "
        "Not a hit. First typed example: adi-wireless-bms-deploy on "
        "ops. Gold China/H20 stays two trees. Gold file still 84,215 "
        "bytes at the 27 August stamp.",
    )
    add_h(doc, "3. Metrics that are not keep rates")
    add_body(
        doc,
        "Transparency (sidecar 2026-09-01T15:00:00+00:00) is a due-clock "
        "slip ledger. In-tree only. Cut/abandon is withdrawn. Implicit "
        "neglect only after a dated clock. Someday wants with no clock "
        "are not ignore. Gold: 6 of 15 dated trees unanswered after "
        "the clock (40%). First-seed books then read 12/12 and 19/19 "
        "because clocks were due and later calls were not walked — "
        "coverage, not a finished trust ranking. Not desk_trust.",
    )
    add_body(
        doc,
        "Quant / expire (sidecar 2026-09-01T16:00:00+00:00). Stated "
        "horizon is seed.clock. Walk does not invent a clock, expire, "
        "or quant binding. expired closes and does not score. Silence "
        "alone is never a miss. Local narrative_quant.parquet only. "
        "Gold never gets a quant apply. First ops examples: "
        "ibm-promontory-watson qualitative expire FY2018-Q3 → expired "
        "/ unknown. STRW $150–160M spend stays pending (measure does "
        "not bind the REIT bogey).",
    )
    add_h(doc, "4. Management regimes")
    add_body(
        doc,
        "Roz page Management Regimes. CEO catalog started at 72 "
        "entries on 43 names, now 74 after AAPL Cook→Ternus "
        "(FY2026-Q4, planned) and APH R. Adam Norwitt (FY2009-Q1, "
        "current). Four views: timeline, comparison, transfer ledger, "
        "multi-company. Transfer kinds: single_regime, prior_closed, "
        "inherited_adopted, inherited_closed_by_successor, "
        "inherited_overdue, inherited_ignored. Dual-surface: Roz and "
        "workshop HTML stay identical.",
    )
    add_h(doc, "5. The pull became a machine (the wow)")
    add_body(
        doc,
        "1 September: one LLM pass per ticker over missed cue rows. "
        "43 tickers, 1,143 cue rows → 310 proposed seeds. Terminal "
        "scoring on the then-open trees: 41 → 21 actionable (7 "
        "delivered, 3 hit, 2 missed, 9 expired, 20 none). Propose, "
        "never auto-insert. Excerpt provenance flagged: 12 / 310 "
        "seeds were not verbatim.",
        ink=True,
    )
    add_body(
        doc,
        "3 September: review workbook, 351 rows (41 verdicts + 310 "
        "seeds) across 43 ticker tabs. Blanket review: high-confidence "
        "seeds accepted (user_blanket_2026-09-03), mediums → "
        "provisional, 28 verbatim verdicts → provisional, non-verbatim "
        "→ Re-cite. Desk Autopilot then inserts what clears a "
        "deterministic bar (verbatim excerpt, validates, not a "
        "duplicate sentence). Budget cap $1 / ticker. Kill switch "
        "DESK_AUTOPILOT=0. NVIDIA gold stays hands-off.",
    )
    add_body(
        doc,
        "Roz fires autopilot after every post_call walk and at the "
        "end of onboard. As of 3 September after AVGO and the "
        "append_node idempotency fix: ops overlay 207 trees (155 "
        "confirmed, 52 provisional), healthcare 187 trees (141 "
        "confirmed, 46 provisional). That is the 44-skeleton book "
        "the room saw on 31 August, walked. Do not call provisional "
        "trees gold-scored. Do not quote those overlays as a keep "
        "rate.",
        ink=True,
    )
    add_h(doc, "6. The monitor was never watching")
    add_body(
        doc,
        "Forensic read of the live monitor DB, 2 September. Eight "
        "events in five weeks. Seven of eight were manual_override=1. "
        "OPAL is the only event that ever armed automatically. The "
        "watched set was 5 of 45 onboarded names. Eligibility and "
        "the research book had been the same list. They are not.",
    )
    add_body(
        doc,
        "Monitored universe is now every onboarded company (overlays "
        "union in-code registry). skip_book_sync no longer drops a "
        "name from watch. Measured: 5 watched → 45. Event source is "
        "Quartr MCP calendar publish, not a one-file inbox. Rank IC "
        "book stays the comparison set. ADSK’s 27 August call was "
        "never processed — that is the honest example, not a Rank IC "
        "miss.",
    )
    add_h(doc, "7. Case study — two Independents, full Roz process")
    add_body(
        doc,
        "Datadog and Lumentum. Independent sector, not XLK, not in "
        "the 17 August Rank IC book. Industry tag is tech. Overlay "
        "file is desk_hc_v2 — a filename, not a healthcare claim. "
        "DDOG is the younger book (first call 12 Nov 2019, latest "
        "FY2026-Q2, 6 Aug). LITE is the ten-year target (June-30 "
        "fiscal, first call 1 Nov 2018, latest FY2026-Q4, 11 Aug). "
        "CRWV stays on the desk as the precedent, not the hero.",
    )
    add_body(
        doc,
        "What actually runs on a new name. (1) Pull — Quartr MCP "
        "only: every FY earnings call and every conference, Q&A "
        "tails included. (2) Score — Structured Narrative, eight "
        "dimensions, four stages: level this quarter, delta vs last "
        "call, surprise vs Street, novelty on the three narrative "
        "dims. (3) Join — feature panel + quant_z_pit, the print "
        "versus this company’s own surprise history as of the call "
        "date. (4) Brief — Post-Call Brief, five questions. "
        "(5) Remember — Claims Desk trees, verbatim clocks. "
        "(6) This meeting — call scorecard; delivery among aged "
        "trees only. (7) Rank IC stays parked. Conferences seed "
        "clocks. FY restates or ignores them.",
        ink=True,
    )
    add_h(doc, "8. Case study — outcomes, the book that landed")
    add_body(
        doc,
        "DDOG: 28 FY (FY2019-Q3–FY2026-Q2), 37 conferences, 25/28 "
        "four-stage (FY2021-Q1 and FY2025-Q3 missing novelty; "
        "FY2019-Q3 prior-only), 200 panel rows, 65 scorecard rows, "
        "10 overlay trees. LITE: 31 FY (FY2019-Q1–FY2026-Q4), 26 "
        "conferences, 232 panel rows, 14 trees. Do not invent 2016. "
        "Quartr has nothing earlier. Delivery on both first-seed "
        "books is an em dash — 10 silent / 10 never-touched on "
        "DDOG FY2026-Q2, 14 / 14 on LITE FY2026-Q4. Inventory is "
        "not a keep rate.",
        ink=True,
    )
    add_h(doc, "9. Case study — meaning")
    add_body(
        doc,
        "Datadog FY2026-Q2 is the gap. Demand +1.9 / surprise +1.2 / "
        "z −0.29 / gap 1.49. They sold 36% growth and record "
        "sequential adds. The call is ahead of the print. Guidance "
        "level still +1.2 (dollars to $4.45–$4.47B), delta −0.8. "
        "Q3 growth 28–29% from 36% after they de-risked the largest "
        "customer. Confidence still +1.6. Novelty −0.8 is the new "
        "caveat, not a tone collapse. On the hook: FedRAMP High "
        "(pipeline not converting), Bits SRE now GA, 30% R&D, "
        "India/Brazil due this quarter.",
        ink=True,
    )
    add_body(
        doc,
        "Lumentum FY2026-Q4 is the split. $1.01B, $1.25B guided. "
        "Demand +2.0 / z +0.15 / agree — confirmation, not a hidden "
        "gap. Margins +2.0 / surprise +1.5 / z −0.38 / gap 1.88 — "
        "crossed 50% gross a year early versus their own $2B model. "
        "Cap alloc z +1.14 agrees (retired $1.1B converts). Macro "
        "−0.5 is the only red cell: Chinese InP just entered the "
        "script. Live clock $2B / 40% in 18–24 months. Dead clock "
        "$600M / 17–20% — already behind $1.01B / ~40%.",
        ink=True,
    )
    add_body(
        doc,
        "Roz does not stamp “bullish.” Divergence (agree False + a "
        "large gap) is the use case. Slope vs adjective. Novelty is "
        "an object detector. A tree remembers what a quarter rank "
        "forgets. On XLK — not these names — raw tone reversed. "
        "Narrative–quant agreement was flat three weeks, ~0.10 by "
        "weeks 6–9, ~0.17 by ~day 90. These two show the machinery. "
        "They are not a live IC. Leave-behind: "
        "Roz_Independent_Case_Study.html.",
    )
    add_h(doc, "10. Call Scorecard is now a live visual")
    add_body(
        doc,
        "The old scatter dropped every row without a delivery rate. "
        "For CRWV that hid the name. Company Timeline now plots "
        "every call in calendar order. Y-axis is transparency. "
        "Color = earnings vs conference. Triangle = open book. Size "
        "= open goals / new seeds. Hover for the event. Click a "
        "point to pin goal health. Delivery stays a tooltip em dash "
        "until a promise settles — never plotted as zero. Period "
        "Snapshot is a transparency bar for every name at that call. "
        "The delivery-vs-transparency quadrant remains, only for "
        "settled points. Sidecar 1,712 rows. Demo: Independent → "
        "DDOG or LITE → Company Timeline. CRWV still plots.",
    )
    add_h(doc, "11. What this is for — the dead window")
    add_body(
        doc,
        "The five-day window between the call and consensus revisions "
        "is the job. Post-Call Brief answers five questions on one "
        "screen: did they deliver, are forward commitments stronger "
        "or weaker, what resolved, what is still open or silent, "
        "and who is management. Claims Desk holds the trees. The "
        "scorecard is the trajectory. Rank IC still ranks the "
        "quarter. They are different objects.",
    )
    add_h(doc, "12. Still true / still not done")
    add_body(
        doc,
        "Gold 6 / 7 and Flex 1 / 1 unchanged. LLY dividend and MSFT "
        "BUILD still open. Depth cohort is started, not closed. "
        "Healthcare is still not a Rank IC. production_v1 is frozen. "
        "Rank IC HTML is stale vs CRWV — out of scope for this talk. "
        "Provisional overlay trees are not gold. CRWV, DDOG, and "
        "LITE delivery is not 0%. Autopilot does not invent "
        "delivered / hit / missed.",
    )
    add_h(doc, "13. Do not say")
    add_body(
        doc,
        "Ops or healthcare keep-rate 0%. CRWV / DDOG / LITE delivery "
        "0%. Healthcare Rank IC. “All twenty healthcare names are "
        "scored gold.” Provisional means confirmed gold. Kept means "
        "delivered. Horizon slip-miss means the tree is missed. Path "
        "ID as the verdict. The remaining 125. A new LLM. "
        "transcripts_raw. Mixing books. Promotion of production_v1. "
        "Depth already done. Flex undone. 45/45 fully live "
        "unattended cycle. Invented delivery rates.",
    )
    add_h(doc, "14. Order")
    add_body(
        doc,
        "Same Rank IC book → 31 August recap (tree, three books, "
        "gold 6 / 7) → afternoon after the talk (honest desk, cull, "
        "ISRG / CRM) → transparency and expire → regimes → "
        "autopilot wow (44 → hundreds) → monitor was manual → "
        "case study (process, book, Datadog gap, Lumentum split, "
        "meaning) → live scorecard → Post-Call Brief as the dead "
        "window → questions. Demo last, not first. Leave-behind is "
        "the HTML.",
    )
    dest = DESK / "Roz_Weekly_Presentation_Outline_2026-09-09.docx"
    doc.save(dest)
    return dest


def write_script() -> Path:
    doc = Document()
    style_doc(doc)
    add_title(doc, "Roz weekly script — 9 September 2026")
    add_body(
        doc,
        "Read this. Do not invent a healthcare Rank IC. Do not call "
        "provisional trees gold. Do not give CRWV, DDOG, or LITE a "
        "delivery rate.",
    )
    add_h(doc, "Open")
    add_body(
        doc,
        "Same Rank IC book as 31 August. asof, generated 17 August "
        "2026, 17:28:40 UTC. I am not promoting anything. Last time "
        "we met, around noon on the 31st, I showed you a tree instead "
        "of a row, three books, and NVIDIA gold at six over seven. "
        "Tech and healthcare were first seeds. The leftover dump was "
        "the hole. We said tighten the pull, then go deep on four "
        "names. This week is what happened after that room.",
    )
    add_h(doc, "The object you already know")
    add_body(
        doc,
        "A Rank IC ranks a quarter. A tree starts from one seed cite "
        "and grows edges. Restated, deferred, silent, delivered, "
        "missed. The walker may mark silent. It may not invent a "
        "close. NVIDIA gold is locked. Thirty-three trees. Cue recall "
        "one hundred percent. Deliver six over seven. The miss is "
        "H20 in FY2026 Q3. Flex on the old Claims Desk page is still "
        "kept and delivered. Path ID is still not the verdict.",
    )
    add_h(doc, "The afternoon of the 31st")
    add_body(
        doc,
        "Option A rode along that afternoon. Sector now picks the "
        "book. Healthcare is not empty behind a tech filter. An em "
        "dash is labeled as no scored closes yet, not zero percent.",
    )
    add_body(
        doc,
        "The leftover list went from twelve hundred fifty openings "
        "to eleven hundred forty-five. Ops five hundred fifty-six. "
        "Healthcare five hundred eighty-nine. That is a shorter "
        "worklist. It is not a short one. Gold recall stayed one.",
    )
    add_body(
        doc,
        "We un-paused the four-name depth file. Intuitive da Vinci X "
        "is delivered in FY2017 Q2. Eleven of one hundred sixty-six "
        "systems were X, first clinical use in Germany, inside the "
        "clock. Salesforce over twenty billion is a hit in FY2021 Q4, "
        "not the COVID-year guide. Lilly’s December 2016 dividend is "
        "still open. Microsoft BUILD is still open. I will not close "
        "those from silence.",
    )
    add_h(doc, "Two scores that are not keep rates")
    add_body(
        doc,
        "Transparency asks whether a dated promise was mentioned "
        "after its clock. Gold: six of fifteen dated trees had no "
        "later cite. Forty percent. That is not desk trust. Expire "
        "closes a clock without scoring it. Silence is not a miss. "
        "Quant reads a local parquet. Gold does not get a quant "
        "apply. If the actual is missing at the stop, the tree is "
        "expired unknown, not missed.",
    )
    add_h(doc, "Who is on the hook")
    add_body(
        doc,
        "There is now a Management Regimes page. CEO catalog, "
        "seventy-four regimes. Apple handed from Cook to Ternus in "
        "FY2026 Q4, planned. Amphenol is Norwitt from FY2009 Q1. "
        "You can see inherited promises the successor closed, and "
        "inherited promises they ignored. Roz and the workshop HTML "
        "show the same numbers.",
    )
    add_h(doc, "The pull")
    add_body(
        doc,
        "This is the week. On the first we ran one model pass per "
        "name over the leftover cues. Forty-three tickers. Three "
        "hundred ten proposed seeds. Then a close pass on the open "
        "trees: twenty-one actionable verdicts. Propose only. Twelve "
        "of those three hundred ten quotes were not verbatim. They "
        "do not get in.",
    )
    add_body(
        doc,
        "On the third we reviewed all of them in one workbook. Three "
        "hundred fifty-one rows. High-confidence seeds were accepted. "
        "Mediums went provisional. Verbatim verdicts went provisional. "
        "Paraphrases went to re-cite. Then the autopilot was allowed "
        "to insert what clears a hard bar: verbatim, valid tree, not "
        "a duplicate sentence. One dollar per ticker. Gold is off "
        "limits. Roz now runs that loop when a call lands and when "
        "a name is onboarded.",
    )
    add_body(
        doc,
        "The book you saw on the 31st had forty-four skeleton trees. "
        "After AVGO and after we stopped the overlay from appending "
        "the same node eight times: two hundred seven ops trees, one "
        "hundred eighty-seven healthcare trees. One hundred fifty-five "
        "and one hundred forty-one of those are confirmed. The rest "
        "are provisional. That is not a keep rate. That is a walked "
        "desk.",
    )
    add_h(doc, "The monitor")
    add_body(
        doc,
        "I owe you an honest sentence on automation. In five weeks "
        "the live monitor database had eight events. Seven were "
        "typed in by hand. One name, OPAL, armed itself. We were "
        "watching five companies out of forty-five. Healthcare "
        "onboard had been told to skip the research book, and that "
        "quietly dropped them from watch. The old tech book had no "
        "overlay files, so the intersection dropped AAPL and MSFT "
        "too. Those are now separate questions: who we watch, versus "
        "who Rank IC ranks. Watch is forty-five. Rank IC is still "
        "the 17 August book. Autodesk’s 27 August call never ran. "
        "That is the example.",
    )
    add_h(doc, "Case study — process")
    add_body(
        doc,
        "The live names are Datadog and Lumentum. Independent. Not "
        "in XLK. Not in the 17 August Rank IC book. They sit on the "
        "healthcare overlay because that is the file, not because "
        "they are drug companies. CoreWeave stays on the desk. It "
        "is not the hero.",
    )
    add_body(
        doc,
        "Here is what actually runs on a new name. We pull every "
        "earnings call and every conference through Quartr in "
        "Cursor — Q&A tails included. Structured Narrative scores "
        "eight dimensions on four stages: how this quarter looks, "
        "whether the story got better or worse versus last call, "
        "whether they were more bullish than the Street, and "
        "whether they said something new. That joins a point-in-time "
        "print z — this company’s own surprise history, as of the "
        "call date. Then a brief, then Claims Desk clocks, then a "
        "scorecard row for this meeting. Rank IC stays parked. "
        "Conferences are where a lot of clocks are first said. "
        "FY is where they get restated or ignored.",
    )

    add_h(doc, "Case study — outcomes")
    add_body(
        doc,
        "Datadog: twenty-eight earnings transcripts, thirty-seven "
        "conferences. First call 12 November 2019. Twenty-five of "
        "twenty-eight FY quarters have all four stages. Two hundred "
        "panel rows. Sixty-five scorecard rows. Ten overlay trees. "
        "Latest quarter FY2026 Q2, sixth of August. Lumentum: "
        "thirty-one earnings, twenty-six conferences. First call "
        "1 November 2018. June-thirtieth fiscal. Two hundred "
        "thirty-two panel rows. Fourteen trees. Latest quarter "
        "FY2026 Q4, eleventh of August. Do not invent 2016. Delivery "
        "on both books is a dash. Ten silent and fourteen silent. "
        "Do not say zero percent. Tree count is inventory, not a "
        "keep rate.",
    )

    add_h(doc, "Case study — meaning")
    add_body(
        doc,
        "Datadog just ran a confident call on a slightly negative "
        "print z. Demand is plus 1.9, surprise plus 1.2, z minus "
        "0.29, gap 1.49. They sold thirty-six percent growth and "
        "record sequential adds. The call is ahead of the number. "
        "Guidance is still plus 1.2 on the adjective — they raised "
        "dollars — and minus 0.8 on the slope. Q3 growth twenty-eight "
        "to twenty-nine from thirty-six, after they de-risked the "
        "largest customer. Confidence did not collapse. Novelty "
        "minus 0.8 is the new caveat. FedRAMP High is still open. "
        "Bits SRE went GA. India/Brazil is due this quarter. Ask it.",
    )
    add_body(
        doc,
        "Lumentum is the contrast. A billion-oh-one quarter, one "
        "twenty-five guided. Demand agrees with the print. Cash "
        "agrees — they retired 1.1 billion of converts, z plus 1.14. "
        "Margins do not. Level pinned at plus 2.0, surprise plus 1.5, "
        "z still minus 0.38, gap 1.88. They crossed fifty percent "
        "gross a year early versus their own two-billion model. "
        "Macro is the only red cell, minus 0.5. Chinese indium "
        "phosphide just entered the script. The six-hundred-million "
        "/ seventeen-to-twenty clock is already dead. The live clock "
        "is two billion / forty percent in eighteen to twenty-four "
        "months.",
    )
    add_body(
        doc,
        "Roz does not stamp bullish. Divergence is the use case — "
        "the call saying something the print z does not. Slope "
        "versus adjective. Novelty is an object detector, not a "
        "louder tone. A tree remembers what a quarter rank forgets. "
        "On XLK, not these names, raw tone reversed. Agreement was "
        "flat three weeks and built later. These two show the "
        "machinery on a live print. They are not a live IC.",
    )
    add_h(doc, "The chart")
    add_body(
        doc,
        "Until last night the scorecard chart threw away any call "
        "without a delivery rate. CoreWeave vanished. It now plots "
        "every tracked call, in time. Transparency on the vertical. "
        "Earnings in one color, conferences in the other. Size is "
        "open goals or new seeds. You can click a point. Delivery "
        "is a dash in the tooltip until a promise settles. We do "
        "not draw that dash as zero. If you want the old quadrant, "
        "it is still there for names that have a settled rate.",
    )
    add_h(doc, "Close")
    add_body(
        doc,
        "If you want to see it: the HTML leave-behind first — "
        "Roz_Independent_Case_Study.html, tabs The pair, Datadog, "
        "Lumentum. Then Independent, DDOG or LITE, Claims Desk, "
        "Company Timeline, FY Post-Call Brief. Switch DDOG off the "
        "Citi TMT conference onto FY2026 Q2. LITE latest FY is "
        "FY2026 Q4. Gold is still on Claims Trees, Book NVIDIA, if "
        "you want the 31 August bar. I have not promoted the Rank "
        "IC book. Questions.",
    )
    dest = DESK / "Roz_Weekly_Presentation_Script_2026-09-09.docx"
    doc.save(dest)
    return dest


def write_card() -> Path:
    doc = Document()
    style_doc(doc)
    add_title(doc, "Claims Desk — talking card — 9 September 2026")
    add_body(
        doc,
        "v1 locked book generated_at=2026-08-17T17:28:40+00:00, split "
        "asof-path-id-17aug-book-v1. Demo v1: http://localhost:8501/Claims_Desk. "
        "v2: http://localhost:8501 → Claims Trees. "
        "Regimes: /Management_Regimes. Brief: /Post_Call_Brief. "
        "Gold stamp 2026-08-27T18:02:00+00:00. "
        "Ops stamp 2026-08-31T14:18:00+00:00. "
        "Healthcare stamp 2026-08-31T14:45:00+00:00. "
        "Scorecard generated_at=2026-09-09T01:27:26Z, n=1712.",
    )
    add_h(doc, "Say these definitions once")
    add_body(
        doc,
        "Tree — one seed cite, then edges. Walk may add restated or "
        "silent. Walk does not invent delivered, hit, or missed.",
        ink=True,
    )
    add_body(
        doc,
        "Kept (v1) — the same object came back. Not delivered.",
        ink=True,
    )
    add_body(
        doc,
        "Slipped — silent plus a due clock. Not missed on the tree.",
        ink=True,
    )
    add_body(
        doc,
        "Deliver / hit — scored terminals only. Unresolved is an em "
        "dash, not 0%. Provisional overlay ≠ gold.",
        ink=True,
    )
    add_body(
        doc,
        "Transparency — due-clock slip, in-tree only. Not desk_trust. "
        "Expire — closes, does not score. Silence ≠ miss.",
        ink=True,
    )
    add_h(doc, "Quote these — still the 31 August bars")
    add_body(
        doc,
        "v1 Flex: ADSK FY2022-Q2 launch, FY2022-Q3 live usage mix. "
        "Kept and delivered. Path ID miss. Book deliver 1 / 1.",
    )
    add_body(doc, "Seed — ADSK FY2022-Q2 · novelty_view · competitive_position")
    add_quote(
        doc,
        "At the end of September, we will launch a new pay-as-you-go "
        "consumption model, called Flex. It matches the customer's cost "
        "with their usage.",
    )
    add_body(doc, "Follow-up — ADSK FY2022-Q3")
    add_quote(
        doc,
        "One of the things we're seeing with Flex is exactly what we "
        "expected to see. We're seeing a large percent of Flex business "
        "coming in is net new.",
    )
    add_body(
        doc,
        "NVIDIA gold: 33 trees (24 promises, 9 goals). Cue recall 1.0 "
        "on 18 / 18. Deliver 6 / 7 (H20 miss FY2026-Q3). Hit 1 / 1. "
        "Window FY2022-Q2–FY2027-Q1. Full-history deliver 9 / 10. "
        "Gold bytes 84,215 at the 27 August stamp.",
    )
    add_h(doc, "Quote these — since noon 31 August")
    add_body(
        doc,
        "Leftover cull, 31 August afternoon: 1,250 → 1,145 openings "
        "(ops 556, HC 589). Gold recall still 1.0.",
    )
    add_body(
        doc,
        "ISRG da Vinci X delivered FY2017-Q2. 11 of 166 systems were "
        "X; first clinical use in Germany; inside FY2017-Q4 clock. "
        "HC scored deliver 1 / 1 on that tree.",
    )
    add_body(
        doc,
        "CRM $20B next goal hit FY2021-Q4 (“over $20 billion in "
        "revenue”). FY2021-Q1 guide is not the hit. Ops scored hit "
        "1 / 1. LLY dividend and MSFT BUILD still open.",
    )
    add_body(
        doc,
        "1 Sep seed batch: 43 tickers, 1,143 cues → 310 proposed "
        "seeds. 12 / 310 not verbatim. Terminal pass: 21 actionable "
        "(7 delivered, 3 hit, 2 missed, 9 expired, 20 none).",
    )
    add_body(
        doc,
        "3 Sep after AVGO + idempotent append: ops 207 trees (155 "
        "confirmed, 52 provisional), HC 187 (141 confirmed, 46 "
        "provisional). Was 44 skeletons on 31 August. Not a keep "
        "rate.",
    )
    add_body(
        doc,
        "Monitor forensics 2 Sep: 8 events / 5 weeks, 7 "
        "manual_override=1. Watched 5 → 45. ADSK 27 Aug call never "
        "ran.",
    )
    add_body(
        doc,
        "CRWV: Independent, desk_hc_v2, Intrator. 19 scorecard rows "
        "(6 FY + 13 CONF). Delivery null on 19 / 19. New seeds 14. "
        "GS 2026 (2026-09-08): 14 open, 14 never touched. 19 briefs. "
        "Deutsche Bank 2025 transparency 0.30, 4 new seeds.",
    )
    add_h(doc, "Demo path")
    add_body(
        doc,
        "localhost:8501 → Independent → CRWV → Claims Desk → "
        "Company Timeline (19 points). Then Post Call Brief → CRWV "
        "→ Goldman Sachs 2026. Optional: Claims Trees → nvda_gold_v2 "
        "for 6 / 7. Do not open Rank IC HTML as if CRWV is in it.",
    )
    add_h(doc, "Do not say")
    add_body(
        doc,
        "Ops or healthcare keep-rate 0%. CRWV delivery 0%. Healthcare "
        "Rank IC. Provisional is gold. Kept means delivered. Horizon "
        "slip means the tree missed. Path ID as the verdict. The "
        "remaining 125. Mixing books. production_v1. Depth done. "
        "Flex undone. Unattended 45-name cycle. Invented rates.",
    )
    dest = DESK / "Roz_Claims_Desk_Talking_Card_2026-09-09.docx"
    doc.save(dest)
    return dest


def main() -> int:
    DESK.mkdir(parents=True, exist_ok=True)
    paths = [write_outline(), write_script(), write_card()]
    for path in paths:
        print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
