"""Talk one-pagers — sell the latest-call read, not a metric tutorial."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DESK = Path(
    r"C:\Users\BobbyWhittaker\OneDrive - Cassius Capital\Desktop"
    r"\Research Presentations"
)
REPO_OUT = ROOT / "data" / "case_study_pull"

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x55, 0x55, 0x55)


def _style(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    pf = style.paragraph_format
    pf.space_after = Pt(3)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE


def _h(doc: Document, text: str, size: int = 13) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = NAVY
    run.font.name = "Calibri"


def _p(doc: Document, text: str, *, muted: bool = False) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    run.font.color.rgb = MUTED if muted else INK


def _bullet(doc: Document, lead: str, body: str) -> None:
    para = doc.add_paragraph(style="List Bullet")
    para.clear()
    r1 = para.add_run(lead)
    r1.bold = True
    r1.font.name = "Calibri"
    r1.font.size = Pt(10.5)
    r1.font.color.rgb = NAVY
    r2 = para.add_run(body)
    r2.font.name = "Calibri"
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = INK


def write_ddog() -> Path:
    doc = Document()
    _style(doc)
    _h(doc, "Datadog FY2026-Q2 — the call is ahead of the print", 16)
    _p(
        doc,
        "Call 6 Aug 2026 · scored against DDOG’s own history (28 FY + 37 "
        "conferences) · Independent, not XLK / Rank IC",
        muted=True,
    )
    _p(
        doc,
        "Datadog just ran a confident call on a slightly negative print z. "
        "The tape they sold is stronger than their own surprise history. The "
        "only thing that got worse versus last quarter is the slope: Q3 "
        "growth guided to 28–29% from 36% after they de-risked the largest "
        "customer. That is the position — not “Datadog is bullish.”",
    )

    _h(doc, "This call vs the print vs last quarter")
    _bullet(
        doc,
        "Demand  +1.9  /  surprise +1.2  /  z −0.29  /  gap 1.49.  ",
        "Surprise says they beat the Street’s demand tape. The print z says "
        "this beat is ordinary versus their own history. They sold 36% "
        "growth, record sequential adds, non-AI high-20s, new logos 25% → "
        "30%. The call is running ahead of the number. Do not stop at the z.",
    )
    _bullet(
        doc,
        "Margins  +0.6  /  in_line +0.2  /  z −0.64  /  gap 0.84.  ",
        "Surprise is flat — they did not sell a margin beat. The z is the "
        "softest print on the page (gross 79.6% vs 80.7%). Level stays "
        "modest because they flagged reinvestment. Op. margin 23% (+1 pt "
        "seq) is cleaner than the z. They are not overclaiming margins. "
        "The print is the soft spot.",
    )
    _bullet(
        doc,
        "Earnings power  +1.5  /  surprise +0.8  /  z −0.61  /  gap 1.41.  ",
        "Surprise says EPS was better than the Street (+9.3%; FY $2.50–$2.54 "
        "vs $2.45). The z still says this is not a rare print versus their "
        "history. Same as demand: more bottom-line power than the "
        "distribution credits, after eating the large-customer headwind.",
    )
    _bullet(
        doc,
        "Capital allocation  +0.7  /  surprise +1.0  /  z −0.51  /  gap 1.51.  ",
        "Surprise is the cash: FCF $279M vs $212M, capex $9.9M vs $49.5M. "
        "Level is only +0.7 — no new buyback story. The beat is real. They "
        "refused to turn it into a promise. Delta flat vs last call.",
    )
    _bullet(
        doc,
        "Guidance  +1.2  /  delta −0.8  /  surprise −0.3  /  agree.  ",
        "This is the tell. Level still positive — dollars raised to "
        "$4.45–$4.47B. Delta says the slope broke versus last quarter. "
        "Surprise is slightly bearish vs Street because Q3 28–29% from 36% "
        "is the new object. The one place words and print agree is the "
        "deceleration. Trade that, not “raised guide.”",
    )
    _bullet(
        doc,
        "Confidence  +1.6  /  delta −0.2  /  novelty −0.8.  ",
        "They still sound like Datadog (“booming”, “record”). Novelty is "
        "the watch item: first time the largest-customer usage cut is in "
        "the script. Confidence did not collapse. A new caveat entered "
        "the book.",
    )
    _bullet(
        doc,
        "Competitive  +1.7  /  delta +0.7  /  novelty +1.2.  ",
        "New objects, not louder adjectives: 750 AI customers including all "
        "top-10, MCP tool-calls 22× vs Q4’25, Bits Security Analyst off "
        "SIEM. Versus last call, the product tape improved.",
    )
    _bullet(
        doc,
        "Macro  +0.3  /  flat  /  novelty +0.2.  ",
        "Dead air. FedRAMP still “building, not converting.” Same script "
        "as last quarter. The FedRAMP High tree is still open.",
    )

    _h(doc, "What they are on the hook for")
    _p(
        doc,
        "10 trees. FY2026-Q2 scorecard: 10 silent, 10 never-touched. "
        "Delivery is an em dash — starter book, not a keep rate.",
        muted=True,
    )
    _bullet(
        doc,
        "FedRAMP High — ",
        "“We're going for FedRAMP high… Department of Defense and others.” "
        "CONF-2025-11-18. This call: pipeline still not converting.",
    )
    _bullet(
        doc,
        "Bits SRE GA — ",
        "“Today, we went GA on our Bits SRE.” CONF-2025-12-02. Earlier: "
        "“still in private beta” (CONF-2025-11-18). The book just closed a beta.",
    )
    _bullet(
        doc,
        "Flex Logs + Cloud SIEM 2025 — ",
        "“scale quite a bit” / “more aggressive in selling Cloud SIEM.” "
        "Clocks FY2025-Q4. Due, not terminal-scored.",
    )
    _bullet(
        doc,
        "30% of revenue in R&D — ",
        "“We invest around 30% of our top line in engineering, and we keep "
        "that going.” The capital-allocation clock they actually own.",
    )
    _bullet(
        doc,
        "25%+ margin — ",
        "“25% plus… cash flow margins 200-300 bps above that.” Q2 printed "
        "23% / 25% FCF. Goal, not a dated quarter.",
    )
    _bullet(
        doc,
        "India / Brazil 50–100, then double — ",
        "Clock FY2026-Q2. Due this quarter. Scorecard has not marked it. Ask it.",
    )
    _bullet(
        doc,
        "Core observability 5–10× — ",
        "CONF-2024-03-05. Long-horizon. Still open.",
    )

    dest = DESK / "Roz_DDOG_FY2026Q2_OnePager.docx"
    dest2 = REPO_OUT / "Roz_DDOG_FY2026Q2_OnePager.docx"
    DESK.mkdir(parents=True, exist_ok=True)
    REPO_OUT.mkdir(parents=True, exist_ok=True)
    doc.save(dest)
    doc.save(dest2)
    return dest


def write_lite() -> Path:
    doc = Document()
    _style(doc)
    _h(doc, "Lumentum FY2026-Q4 — the print agrees, except where it doesn’t", 16)
    _p(
        doc,
        "Call 11 Aug 2026 · June-30 fiscal · scored against LITE’s own "
        "history (31 FY + 26 conferences) · Independent, not XLK / Rank IC",
        muted=True,
    )
    _p(
        doc,
        "Lumentum just printed $1.01B and guided $1.25B. Demand, guidance, "
        "and the balance sheet agree with the number. Margins and EPS do "
        "not — they are talking a +2.0 operating story against a still-"
        "negative PIT z. Versus last call they stepped up again, and they "
        "finally put Chinese InP on the tape. The old $600M / 17–20% clock "
        "is already dead. The live clock is $2B / 40%.",
    )

    _h(doc, "This call vs the print vs last quarter")
    _bullet(
        doc,
        "Demand  +2.0  /  surprise +1.4  /  z +0.15  /  agree.  ",
        "Surprise says they blew past the Street ($1.01B, +109% YoY, Q1 "
        "$1.25B vs $1.165B). The z agrees on sign. They cannot service all "
        "customer demand. Confirmation of the squeeze — not a hidden gap.",
    )
    _bullet(
        doc,
        "Margins  +2.0  /  surprise +1.5  /  z −0.38  /  gap 1.88.  ",
        "Surprise says they sold a bigger margin tape than the Street "
        "(gross 50.4% vs 48.1%; Q1 op. 39.5–40.5%). The z is still "
        "negative versus their own history. They crossed 50% a year early "
        "versus the old $2B model. Words ahead of the historical "
        "distribution — same Datadog pattern, on a name that just printed "
        "2,160 bps of YoY op. expansion.",
    )
    _bullet(
        doc,
        "Earnings power  +2.0  /  surprise +1.2  /  z −0.40  /  gap 1.60.  ",
        "EPS $3.23 vs $2.97; Q1 $4.05–$4.35 vs $3.65. Surprise: the Street "
        "was behind. The z: this is not rare in their recent book. They "
        "beat their own $2.85–$3.05 guide. The call is pricing in more "
        "leverage than the PIT z.",
    )
    _bullet(
        doc,
        "Capital allocation  +1.5  /  surprise +1.3  /  z +1.14  /  agree.  ",
        "The rare rhyme. Retired $1.1B converts. Net debt −$1.1B vs ~+$1.9B "
        "consensus. CFO $363M vs $212M. Surprise and the print are the "
        "same sentence. Gap 0.16. Do not treat this like the margin gap.",
    )
    _bullet(
        doc,
        "Guidance  +2.0  /  delta +1.8  /  surprise +1.6  /  rev z +0.87.  ",
        "They pulled $1.25B forward more than a quarter and said the whole "
        "target stack gets raised at the next OFC. Versus last call this "
        "is a reset, not a beat. Surprise and the revision z agree. The "
        "Street was late.",
    )
    _bullet(
        doc,
        "Confidence  +1.9  /  delta +0.8  /  novelty +1.2.  ",
        "They beat their own clocks and said so. Novelty is the schedule "
        "pull-forward, not a new adjective. Tone stepped up from an "
        "already-high Q3.",
    )
    _bullet(
        doc,
        "Competitive  +1.9  /  delta +1.0  /  novelty +1.8.  ",
        "New objects: first ELS PO, 70–80% pump-laser share (first time "
        "quantified), NPO as additive TAM, first triple-digit OCS quarter "
        "guided, AXT InP deal. A product-tape rewrite versus last quarter.",
    )
    _bullet(
        doc,
        "Macro  −0.5  /  delta −0.5  /  novelty +0.8.  ",
        "The only red cell. Last quarter this dim was ignored. This quarter "
        "Chinese InP fab competition is on the tape. Political constraints "
        "are a pump-laser tailwind and a geopolitical risk. Do not skip it "
        "because everything else is +2.",
    )

    _h(doc, "What they are on the hook for")
    _p(
        doc,
        "14 trees. FY2026-Q4 scorecard: 14 silent, 14 never-touched, 10 "
        "open-stale. Delivery is an em dash. The superseded clocks are the "
        "point — you can see the book aging in public.",
        muted=True,
    )
    _bullet(
        doc,
        "$2B / 40% in 18–24 months — ",
        "“exit at $2 billion and 40% operating margins anywhere from 18 "
        "months-24 months from now.” FY2026-Q3, clock FY2027-Q4. Live.",
    )
    _bullet(
        doc,
        "CPO $100M Q4 CY2026 — ",
        "“$10 million of ROC, $10 million of OCS in Q1 going to $100 million "
        "of incremental dollars for us in Q4.” CONF-2025-12-08.",
    )
    _bullet(
        doc,
        "$400M CPO backlog H2 CY2026 — ",
        "“on track to be shipping this $400 million of backlog in the H2 of "
        "this calendar year.” FY2026-Q3.",
    )
    _bullet(
        doc,
        "$100M quarterly OCS — ",
        "Roadmap $10M → $100M. This call guided the first triple-digit OCS "
        "quarter. Coming due.",
    )
    _bullet(
        doc,
        "Greensboro fab 2028 — ",
        "“shipping out of this Greensboro fab by 2028.” Capacity for the $2B exit.",
    )
    _bullet(
        doc,
        "$600M / 17–20% — already dead.  ",
        "Q4 printed $1.01B and guided 39.5–40.5%. Keep it so the room sees "
        "the book compounding. $500M-by-CY2025 is the same ancestor.",
    )
    _bullet(
        doc,
        "InP / Thailand / 200G EML / NeoPhotonics synergies — ",
        "Capacity and synergy ancestors. The constraint behind “we cannot "
        "service all customer demand” is the same tree, three years later.",
    )

    dest = DESK / "Roz_LITE_FY2026Q4_OnePager.docx"
    dest2 = REPO_OUT / "Roz_LITE_FY2026Q4_OnePager.docx"
    doc.save(dest)
    doc.save(dest2)
    return dest


def main() -> None:
    print(write_ddog())
    print(write_lite())


if __name__ == "__main__":
    main()
