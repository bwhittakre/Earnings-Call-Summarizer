"""Write 27 Aug weekly outline, script, and Claims Desk card to the Desktop."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt, RGBColor

DESK = Path(r"C:\Users\BobbyWhittaker\OneDrive - Cassius Capital\Desktop")
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
MUTED = RGBColor(0x44, 0x44, 0x44)
INK = RGBColor(0x1A, 0x1A, 0x1A)


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = MUTED
    pf = style.paragraph_format
    pf.space_after = Pt(8)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = NAVY
    run.font.name = "Calibri"


def add_h(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = NAVY
    run.font.name = "Calibri"


def add_body(doc: Document, text: str, *, ink: bool = False) -> None:
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.font.color.rgb = INK if ink else MUTED


def add_quote(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = INK


def write_outline() -> Path:
    doc = Document()
    style_doc(doc)
    add_title(doc, "Roz weekly outline — 27 August 2026")
    add_body(
        doc,
        "Coverage: 26 August talk through this morning. Book for every "
        "Rank IC sentence: asof, generated_at=2026-08-17T17:28:40+00:00. "
        "Claims desk split: asof-path-id-17aug-book-v1, same stamp. "
        "Not a promotion. No holdout. Healthcare is a print, not a Rank IC.",
    )
    add_h(doc, "Through-line (~12 minutes)")
    add_body(
        doc,
        "Last week: the software-novelty U is two curves, and 0_56 is not "
        "the typical quarter. This week you add one operational sentence: "
        "we can now track a typed promise with a follow-up cite. Kept is "
        "not delivered. The bar is Autodesk Flex. Book deliver rate is "
        "1 / 1 on the locked book.",
    )
    add_h(doc, "1. The sentence that wows (term structure — still live)")
    add_body(
        doc,
        "The full-sample software-novelty U-shape is an average of two "
        "different curves (expe-term-structure-v1). Early 2016-Q2..2021-Q1: "
        "+0.183 / −0.030 / +0.112. Late 2021-Q2..2026-Q1: +0.070 / +0.083 / "
        "+0.029. 0_56 is a compound return, not the average of those buckets.",
    )
    add_body(
        doc,
        "Typical-quarter test failed (expe-term-structure-v3): 4/20 late "
        "periods have 0_56 above the best bucket (2022-Q3, 2022-Q4, "
        "2025-Q1, 2026-Q1). Bucket-return pairwise Spearman 0.052381. "
        "Novelty vs z-sum of the three bucket returns +0.158035, above "
        "late 0_56 +0.135418. The mean premium is noise cancellation.",
    )
    add_h(doc, "2. The new object — Claims Desk")
    add_body(
        doc,
        "One row is one object: type, clock, state, cite, follow-up cite. "
        "28 rows on the locked 17 Aug book (15 pilot + 13 same-beat backfill). "
        "States: kept 9, open 4, subject-changed 15, slipped 0. "
        "Delivery: delivered 1, missed 0, unresolved 3, not-a-promise 24.",
    )
    add_body(
        doc,
        "Kept = the same object came back next quarter. Delivery = a dated "
        "promise actually came true. Deliver rate = delivered / "
        "(delivered + missed). Unresolved clocks and printed facts are an "
        "em dash, not 0%. Book rate 1 / 1 — Autodesk Flex launch only. "
        "Live in Roz at Claims Desk (localhost:8501/Claims_Desk).",
    )
    add_h(doc, "3. The bar — Flex")
    add_body(
        doc,
        "ADSK FY2022-Q2 (calendar 2021-Q2): “we will launch Flex end of "
        "September.” Next quarter FY2022-Q3 is a buried composite, not "
        "the first verbatim. It treats Flex as live usage mix — net new, "
        "occasional, and advanced-product buyers. Verdict: kept and "
        "delivered. Path ID on that row is a miss. Path ID is not the verdict.",
    )
    add_body(
        doc,
        "Three promises still unresolved: Flex transaction-model clock "
        "(fiscal 25/26; silence before the clock is not a miss), "
        "CRM Informatica close, ADSK digital-twin TAM (no next quarter "
        "in the locked book). Do not file the remaining 125 Path ID rows.",
    )
    add_h(doc, "4. Do not say")
    add_body(
        doc,
        "“Novelty is U-shaped” without the era. “0_56 is the average of "
        "the three buckets.” “Kept means they delivered.” “Deliver rate "
        "is 0% for Salesforce.” Healthcare Rank IC. All twenty healthcare "
        "names finished. Promotion, holdout, AI-era. Path ID as the "
        "promise tracker.",
    )
    add_h(doc, "5. Healthcare — honest status")
    add_body(
        doc,
        "Leftover history scored 26 Aug. Tagged pack "
        "narrative_signal_eval_healthcare_large_cap.json, "
        "generated_at=2026-08-26T22:48:11+00:00, 20 names, 253/253 "
        "dimensions. Locked tech book unchanged. production_v1 frozen. "
        "No healthcare Rank IC. Remaining panel holes: CI FY2018-Q4, "
        "TMO FY2017-Q2. LLY is US5324571083 only. Lloyds GB0005163141 "
        "is dropped. UNH IBES ticker is UNIH.",
    )
    add_h(doc, "6. Order")
    add_body(
        doc,
        "Same book / not a promotion → term-structure wow and "
        "typical-quarter fail → Claims Desk object (kept ≠ delivered) → "
        "Flex quotes → three open clocks → healthcare print, not Rank IC → "
        "questions.",
    )
    dest = DESK / "Roz_Weekly_Presentation_Outline_2026-08-27.docx"
    doc.save(dest)
    return dest


def write_script() -> Path:
    doc = Document()
    style_doc(doc)
    add_title(doc, "Roz weekly script — 27 August 2026")
    add_body(doc, "Read this. Do not invent a healthcare Rank IC.")
    add_h(doc, "Open")
    add_body(
        doc,
        "Same book as last week. asof, generated at 17 August 2026, "
        "17:28:40 UTC. I am not promoting anything. I have one new "
        "object on that book, and I want to show you the bar first.",
    )
    add_h(doc, "Wow — still the curves")
    add_body(
        doc,
        "The software-novelty U you have seen on the full sample is an "
        "average of two different curves. Early: a print with a hole in "
        "the middle, plus 18, minus 3, plus 11. Late: a mid-window hump, "
        "plus 7, plus 8, plus 3. Same signal. Opposite relationship to "
        "the combined 0_56 window.",
    )
    add_h(doc, "Then the fail, on purpose")
    add_body(
        doc,
        "Late 0_56 plus 13 and a half looks better than every bucket. "
        "That is not the typical quarter. It is four of twenty: 2022 Q3, "
        "2022 Q4, 2025 Q1, 2026 Q1. The three bucket returns are almost "
        "uncorrelated. Averaging those noisy slice returns recovers more "
        "Rank IC than compounding them. I am not going to present plus 13 "
        "as a horizon that uniformly wins.",
    )
    add_h(doc, "The new object")
    add_body(
        doc,
        "Rank IC ranks a quarter. It does not tell you whether a dated "
        "promise came true. That is now a separate desk. One row is one "
        "object: a type, a clock, a state, a cite, and a follow-up cite. "
        "Twenty-eight rows. Fifteen from the pilot, thirteen same-beat "
        "backfill. Not the remaining 125.",
    )
    add_body(
        doc,
        "Kept means the same object came back. Delivered means the dated "
        "promise actually happened. Those are different. Printed facts "
        "are not promises. A silent clock that is not yet due is open, "
        "not missed. Book deliver rate is delivered over delivered plus "
        "missed. Unresolved does not go in the denominator. On this book "
        "that rate is one over one.",
    )
    add_h(doc, "The bar")
    add_body(
        doc,
        "Autodesk, fiscal 2022 Q2. They said they will launch Flex at the "
        "end of September. Next quarter does not repeat the launch date. "
        "It talks about Flex as live business — net new, occasional, "
        "advanced-product buyers. That is kept, and it is delivered. "
        "The follow-up is a buried composite, not the first line of the "
        "novelty view. Path ID marked that row a miss. Path ID is not "
        "the verdict.",
    )
    add_h(doc, "What is still open")
    add_body(
        doc,
        "Three promises are unresolved. The later Flex transaction-model "
        "clock is fiscal 25 / 26 — silence before the clock is not a miss. "
        "Salesforce has signed to buy Informatica; next quarter does not "
        "say closed or dead. Autodesk’s digital-twin TAM has no next "
        "quarter in the locked book, so it cannot be scored. I will not "
        "file a hundred and twenty-five more Path ID rows to look busy.",
    )
    add_h(doc, "Healthcare")
    add_body(
        doc,
        "Leftover history on the healthcare names is scored and stamped "
        "26 August, 22:48:11 UTC. Twenty names. Two quarters still missing "
        "from those panels: Cigna 2018 Q4, Thermo Fisher 2017 Q2. That is "
        "not a Rank IC. Lilly is the US ISIN only. UnitedHealth’s IBES "
        "ticker is UNIH. The live tech book was not rewritten.",
    )
    add_h(doc, "Close")
    add_body(
        doc,
        "If you want to see it, Claims Desk is in Roz, next to Rank IC. "
        "Flex is open by default. Questions.",
    )
    dest = DESK / "Roz_Weekly_Presentation_Script_2026-08-27.docx"
    doc.save(dest)
    return dest


def write_card() -> Path:
    doc = Document()
    style_doc(doc)
    add_title(doc, "Claims Desk — talking card — 27 August 2026")
    add_body(
        doc,
        "Locked book generated_at=2026-08-17T17:28:40+00:00. "
        "Split asof-path-id-17aug-book-v1. Demo: http://localhost:8501/Claims_Desk",
    )
    add_h(doc, "Say these definitions once")
    add_body(
        doc,
        "Type — printed fact, completed announcement, forward clock, or "
        "pending close. Only the last two are promises.",
        ink=True,
    )
    add_body(
        doc,
        "State — open, kept, slipped, subject-changed. Kept = the object "
        "came back. Slipped is not automatically missed.",
        ink=True,
    )
    add_body(
        doc,
        "Delivery — delivered, missed, unresolved, not-a-promise. "
        "Hand-labeled from the two cites. Not inferred from kept.",
        ink=True,
    )
    add_h(doc, "Book counts (quote these)")
    add_body(
        doc,
        "28 rows. Pilot 15. Same-beat backfill 13. "
        "State: kept 9, open 4, subject-changed 15, slipped 0. "
        "Delivery: delivered 1, missed 0, unresolved 3, not-a-promise 24. "
        "Deliver rate 1 / 1 (100%). Salesforce, Informatica, later Flex, "
        "and digital-twin TAM are em dashes, not zeros.",
    )
    add_h(doc, "Flex — read the two cites")
    add_body(doc, "Seed — ADSK FY2022-Q2 · novelty_view · competitive_position · verbatim")
    add_quote(
        doc,
        "At the end of September, we will launch a new pay-as-you-go "
        "consumption model, called Flex. It matches the customer's cost "
        "with their usage.",
    )
    add_body(doc, "Follow-up — ADSK FY2022-Q3 · novelty_view · competitive_position · composite")
    add_quote(
        doc,
        "One of the things we're seeing with Flex is exactly what we "
        "expected to see. We're seeing a large percent of Flex business "
        "coming in is net new.",
    )
    add_body(
        doc,
        "Coverage: next quarter treated Flex as live usage mix — net new, "
        "occasional, and advanced-product buyers — not a restated launch "
        "date. Verdict: kept and delivered. Path ID hit on this row: false.",
    )
    add_h(doc, "Three unresolved clocks")
    add_body(
        doc,
        "ADSK Flex transaction model — fiscal 25/26. Next quarter silent. "
        "Silence before the clock is not a miss.",
    )
    add_body(
        doc,
        "CRM Informatica — $8 billion pending close. No next-quarter cite "
        "that it closed or died.",
    )
    add_body(
        doc,
        "ADSK digital-twin TAM — no next quarter in the locked book. "
        "Cannot be scored.",
    )
    add_h(doc, "Do not say")
    add_body(
        doc,
        "Kept means delivered. Deliver rate 0% on names with no scored "
        "promise. Path ID hit/miss as the desk verdict. The remaining 125. "
        "A new LLM. transcripts_raw. Healthcare mixed into this book. "
        "Promotion of production_v1.",
    )
    dest = DESK / "Roz_Claims_Desk_Talking_Card_2026-08-27.docx"
    doc.save(dest)
    return dest


def main() -> int:
    DESK.mkdir(parents=True, exist_ok=True)
    paths = [write_outline(), write_script(), write_card()]
    for path in paths:
        print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
