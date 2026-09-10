"""Latest-quarter one-pager for the Independent case study (DDOG / LITE)."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DESK = Path(
    r"C:\Users\BobbyWhittaker\OneDrive - Cassius Capital\Desktop"
    r"\Research Presentations"
)
INV = ROOT / "data" / "case_study_pull" / "inventory.json"
SCORE = ROOT / "data" / "desk_call_scorecard_v1.json"
OVERLAY = ROOT / "data" / "desk_catalog_overlay" / "hc.json"
BRIEFS = ROOT / "data" / "briefs"
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
MUTED = RGBColor(0x44, 0x44, 0x44)
INK = RGBColor(0x1A, 0x1A, 0x1A)

CASES = (
    {
        "ticker": "DDOG",
        "name": "Datadog",
        "period": "FY2026-Q2",
        "isin": "US23804L1035",
        "url": "https://web.quartr.com/companies/6106",
    },
    {
        "ticker": "LITE",
        "name": "Lumentum",
        "period": "FY2026-Q4",
        "isin": "US55024U1097",
        "url": "https://web.quartr.com/companies/6193",
    },
)


def _style(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    pf = style.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE


def _h(doc: Document, text: str, size: int = 14) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = NAVY
    run.font.name = "Calibri"


def _p(doc: Document, text: str) -> None:
    para = doc.add_paragraph(text)
    for run in para.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.font.color.rgb = INK


def _coverage(ticker: str) -> str:
    inv = json.loads(INV.read_text(encoding="utf-8"))
    book = inv[ticker]
    raw = ROOT / "Structured Narrative" / "transcripts_raw"
    n_fy = len(list(raw.glob(f"{ticker}_FY*.txt")))
    conf_dir = ROOT / "data" / "conf_transcripts" / ticker
    n_conf = (
        len([p for p in conf_dir.glob("*.json") if p.name != "manifest.json"])
        if conf_dir.is_dir()
        else 0
    )
    return (
        f"{book['n_earnings']} earnings transcripts available on Quartr "
        f"({book['first_transcript']} through {book['last_transcript']}); "
        f"{n_fy} written on disk. {book['n_conference']} conference / "
        f"investor-day transcripts available; {n_conf} written on disk. "
        f"{book['notes']}"
    )


def _panel_line(ticker: str, period: str) -> str:
    csv_path = (
        ROOT
        / "Structured Narrative"
        / "output"
        / ticker
        / "csv"
        / "feature_panel.csv"
    )
    if not csv_path.is_file():
        return (
            "Feature panel not on disk yet. Narrative / quant for this "
            "quarter are not quoted."
        )
    import csv

    rows = [
        r
        for r in csv.DictReader(csv_path.open(encoding="utf-8"))
        if str(r.get("fiscal_period") or r.get("period") or "") == period
    ]
    if not rows:
        return (
            f"Feature panel exists but {period} is not in it yet. "
            "Do not invent a score."
        )
    bits = []
    for r in rows:
        dim = r.get("dimension") or "?"
        level = r.get("llm_level") or "—"
        qz = r.get("quant_z") or r.get("quant_z_pit") or "—"
        gap = r.get("narrative_quant_gap") or "—"
        bits.append(f"{dim} level={level} quant_z={qz} gap={gap}")
    return f"Panel {period} ({len(rows)} dimensions): " + "; ".join(bits)


def _desk_line(ticker: str, period: str) -> str:
    trees = []
    if OVERLAY.is_file():
        payload = json.loads(OVERLAY.read_text(encoding="utf-8"))
        trees = [
            t
            for t in (payload.get("trees") or [])
            if str(t.get("ticker") or "").upper() == ticker
        ]
    n = len(trees)
    n_prov = sum(
        1
        for t in trees
        if str((t.get("provenance") or {}).get("status") or "") == "provisional"
    )
    n_conf = sum(
        1
        for t in trees
        if str((t.get("provenance") or {}).get("status") or "") == "confirmed"
    )
    entries = []
    if SCORE.is_file():
        entries = [
            e
            for e in (json.loads(SCORE.read_text(encoding="utf-8")).get("entries") or [])
            if str(e.get("ticker") or "").upper() == ticker
        ]
    this = next((e for e in entries if e.get("fiscal_period") == period), None)
    brief = BRIEFS / f"{ticker}_{period.replace('-', '')}_brief.md"
    parts = [
        f"{n} overlay tree(s) on desk_hc_v2 ({n_conf} confirmed, {n_prov} provisional).",
        f"{len(entries)} scorecard row(s).",
    ]
    if this is None:
        parts.append(
            f"No scorecard row for {period} yet. Delivery is an em dash, not 0%."
        )
    else:
        never = this.get("n_never_touched")
        scored = this.get("n_confirmed") or 0
        failed = this.get("n_failed") or 0
        if (scored + failed) == 0:
            parts.append(
                f"{period} has no scored terminals. Delivery is an em dash, "
                f"not 0%. never_touched={never}."
            )
        else:
            parts.append(
                f"{period} delivery_score={this.get('delivery_score')}. "
                f"never_touched={never}."
            )
    if brief.is_file():
        parts.append(f"Post-Call Brief on disk: {brief.name}.")
    else:
        parts.append("Post-Call Brief not written yet.")
    parts.append("Tree count is inventory. It is not a keep rate.")
    return " ".join(str(p) for p in parts)


def main() -> Path:
    DESK.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _style(doc)
    _h(doc, "Roz Independent case study — DDOG and LITE", 18)
    _p(
        doc,
        "10 AM talk one-pager. Both names are Independent, industry_group=tech. "
        "They are not on the locked 17 August Rank IC book. Overlay file is "
        "desk_hc_v2 — a file, not a healthcare claim. FY earnings get narrative "
        "+ quant. Conferences seed Claims Desk only. Do not invent delivery.",
    )
    for case in CASES:
        ticker = case["ticker"]
        period = case["period"]
        _h(doc, f"{ticker} · {case['name']} · {period}")
        _p(doc, f"ISIN {case['isin']}. Quartr: {case['url']}")
        _h(doc, "Coverage", 12)
        _p(doc, _coverage(ticker))
        _h(doc, "Latest quarter", 12)
        _p(doc, _panel_line(ticker, period))
        safe = period.replace("-", "")
        brief = BRIEFS / f"{ticker}_{safe}_brief.md"
        _p(
            doc,
            "Five questions: did they deliver, stronger or weaker, what "
            "resolved, what is still open or silent, who is on the hook. "
            "First-seed books have no scored terminals. Delivery is an em "
            "dash. Do not read a generated 0% as a keep rate.",
        )
        if brief.is_file():
            text = brief.read_text(encoding="utf-8").strip()
            # Drop the auto 0% header so the room copy cannot quote it.
            lines = [
                ln
                for ln in text.splitlines()
                if "Delivery rate:" not in ln and "Transparency score:" not in ln
            ]
            clipped = "\n".join(lines).strip()
            _p(doc, clipped[:2200] + ("…" if len(clipped) > 2200 else ""))
        else:
            _p(doc, "Post-Call Brief not written yet for this period.")
        _h(doc, "Claims Desk", 12)
        _p(doc, _desk_line(ticker, period))
    dest = DESK / "Roz_Independent_Case_Study_DDOG_LITE_2026-09-09.docx"
    doc.save(dest)
    print(dest)
    return dest


if __name__ == "__main__":
    main()
