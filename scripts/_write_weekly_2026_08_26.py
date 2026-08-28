"""Write 26 Aug weekly outline + script to the Desktop."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt, RGBColor

DESK = Path(r"C:\Users\BobbyWhittaker\OneDrive - Cassius Capital\Desktop")
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
MUTED = RGBColor(0x44, 0x44, 0x44)


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = MUTED
    pf = style.paragraph_format
    pf.space_after = Pt(8)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = NAVY
    run.font.name = "Calibri"


def add_h(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = NAVY
    run.font.name = "Calibri"


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.font.color.rgb = MUTED


def write_outline() -> Path:
    doc = Document()
    style_doc(doc)
    add_title(doc, "Roz weekly outline — 26 August 2026")
    add_body(
        doc,
        "Coverage: 19 August talk through this morning. Book for every "
        "Rank IC sentence: asof, generated_at=2026-08-17T17:28:40+00:00. "
        "Not a promotion. No holdout.",
    )
    add_h(doc, "Through-line (~12 minutes)")
    add_body(
        doc,
        "Last week you could defend four sentences on the 17 August book. "
        "This week you add one operational sentence: a newest-first "
        "dimension view will invert a delta. We caught it on UNH. "
        "Healthcare is a print, not a Rank IC.",
    )
    add_h(doc, "1. The sentence that wows")
    add_body(
        doc,
        "The full-sample software-novelty U-shape is an average of two "
        "different curves (expe-term-structure-v1). Early 2016-Q2..2021-Q1: "
        "+0.183 / −0.030 / +0.112. Late 2021-Q2..2026-Q1: +0.070 / +0.083 / "
        "+0.029. 0_56 is a compound, not the average of those buckets.",
    )
    add_body(
        doc,
        "Typical-quarter test failed (expe-term-structure-v3): 4/20 late "
        "periods have 0_56 above the best bucket (2022-Q3, 2022-Q4, "
        "2025-Q1, 2026-Q1). Bucket-return pairwise Spearman 0.052381. "
        "Novelty vs z-sum of the three bucket returns +0.158035, above "
        "0_56 +0.135418. The mean premium is noise cancellation.",
    )
    add_h(doc, "2. Do not say")
    add_body(
        doc,
        "“Novelty is U-shaped” without the era. “0_56 is the average of "
        "the three buckets.” Healthcare Rank IC. All twenty healthcare "
        "names onboarded. Promotion, holdout, AI-era.",
    )
    add_h(doc, "3. Healthcare — honest status")
    add_body(
        doc,
        "Thirteen names have feature panels in a separate research sector, "
        "not mixed into live XLK / SQLite. UNH identity: ISIN US91324P1021, "
        "estpermid 30064860782, Barra USAO6Z1, IBES UNIH. Do not bind UNHC. "
        "Two-quarter test: prior FY2025-Q4, output FY2026-Q1 and FY2026-Q2. "
        "LLY overlay still has Lloyds GB0005163141. Still need sourced "
        "ISINs for DHR, SYK, GILD, VRTX, ELV. No healthcare book stamp.",
    )
    add_h(doc, "4. Order")
    add_body(
        doc,
        "Instrumentation → Lab vs production (frozen) → term-structure wow "
        "and typical-quarter fail → UNH identity and inverted-delta catch → "
        "what is still open.",
    )
    dest = DESK / "Roz_Weekly_Presentation_Outline_2026-08-26.docx"
    doc.save(dest)
    return dest


def write_script() -> Path:
    doc = Document()
    style_doc(doc)
    add_title(doc, "Roz weekly script — 26 August 2026")
    add_body(doc, "Read this. Do not invent a healthcare Rank IC.")
    add_h(doc, "Open")
    add_body(
        doc,
        "Same book as last week. asof, generated at 17 August 2026, "
        "17:28:40 UTC. I am not promoting anything.",
    )
    add_h(doc, "Wow")
    add_body(
        doc,
        "The software-novelty U you have seen on the full sample is an "
        "average of two different curves. Early: a print with a hole in "
        "the middle, plus 18, minus 3, plus 11. Late: a mid-window hump, "
        "plus 7, plus 8, plus 3. Same signal. Opposite relationship to "
        "the combined 0_56 window.",
    )
    add_h(doc, "Then the fail, on purpose")
    add_body(
        doc,
        "Late 0_56 plus 13 and a half looks better than every bucket. "
        "That is not the typical quarter. It is four of twenty: 2022 Q3, "
        "2022 Q4, 2025 Q1, 2026 Q1. The three bucket returns are almost "
        "uncorrelated. Averaging those noisy slice returns recovers more "
        "Rank IC than compounding them. I am not going to present plus 13 "
        "as a horizon that uniformly wins.",
    )
    add_h(doc, "Healthcare")
    add_body(
        doc,
        "Thirteen names are on disk as a separate sector. UnitedHealth "
        "is the first leftover we closed properly. Bloomberg ISIN "
        "US91324P1021. LSEG’s IBES ticker is UNIH, not UNH. If you bind "
        "the prefix you get a different firm. We ran two quarters, not "
        "the full history, and we do not have a healthcare Rank IC.",
    )
    add_h(doc, "The catch")
    add_body(
        doc,
        "The first UNH delta treated Q2 as the prior to Q1 because the "
        "first-write view landed newest-first. Adjacent pairing followed "
        "the list, not fiscal time. That is now sorted before pairing. "
        "I would rather show you a bug we caught than a sector IC we "
        "have not earned.",
    )
    add_h(doc, "Close")
    add_body(
        doc,
        "Lilly’s overlay still has a Lloyds ISIN. Five names still need "
        "a sourced ISIN. The live book was not rewritten. Questions.",
    )
    dest = DESK / "Roz_Weekly_Presentation_Script_2026-08-26.docx"
    doc.save(dest)
    return dest


def main() -> int:
    DESK.mkdir(parents=True, exist_ok=True)
    outline = write_outline()
    script = write_script()
    print(outline)
    print(script)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
