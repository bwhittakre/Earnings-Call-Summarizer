"""Write the 9 Sep final presentation outline and script."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt, RGBColor

DESK = Path(
    r"C:\Users\BobbyWhittaker\OneDrive - Cassius Capital\Desktop"
    r"\Research Presentations"
)
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
MUTED = RGBColor(0x44, 0x44, 0x44)
INK = RGBColor(0x1A, 0x1A, 0x1A)

RANK_IC = "asof, generated_at=2026-08-17T17:28:40+00:00"
GOLD = "2026-08-27T18:02:00+00:00"
OPS = "2026-08-31T14:18:00+00:00"
HC = "2026-08-31T14:45:00+00:00"
SCORECARD = "2026-09-09T01:27:26Z, 1,712 rows"


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = MUTED
    pf = style.paragraph_format
    pf.space_after = Pt(8)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = NAVY
    run.font.name = "Calibri"


def add_h(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = NAVY
    run.font.name = "Calibri"


def add_body(doc: Document, text: str, *, ink: bool = False) -> None:
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.font.color.rgb = INK if ink else MUTED


def add_quote(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = INK


def write_outline() -> Path:
    doc = Document()
    style_doc(doc)
    add_title(doc, "Roz final outline — 9 September 2026")
    add_body(
        doc,
        "Final talk. Two full parts, then a short thank-you. "
        "Part 1 is the summer: how Roz came to be and what it is. "
        "Part 2 is the last week, then the Independent case study — "
        "process, outcomes, meaning. "
        f"Rank IC book is still {RANK_IC}. Not a promotion. "
        f"NVIDIA gold {GOLD}. Ops {OPS}. Healthcare desk {HC}. "
        f"Scorecard {SCORECARD}. No healthcare Rank IC. "
        "No production_v1. No invented delivery rates.",
    )
    add_h(doc, "Through-line (~30 minutes)")
    add_body(
        doc,
        "Part 1 (~15 min). One question all summer: does a cite-backed "
        "read of the call add anything the street has not already "
        "priced? We went from a summarizer, to a scored narrative "
        "joined to a point-in-time quant spine, to Roz — onboard, "
        "score, and then keep the living object. Rank IC ranks a "
        "quarter. Claims Desk tracks the promise. Transparency asks "
        "whether they mentioned it after the clock. The job is the "
        "window before consensus revises.",
    )
    add_body(
        doc,
        "Part 2 (~15 min). Noon 31 August the room asked for B, then C. "
        "This week B ran. 310 proposed seeds, a 351-row review, "
        "Desk Autopilot, 44 skeletons walked to 207 ops / 187 HC. "
        "The monitor had been watching 5 of 45. Then the Independent "
        "case study — process, outcomes, meaning. DDOG FY2026-Q2 is "
        "the gap (words ahead of the print). LITE FY2026-Q4 is the "
        "split (demand agrees; margins do not). Delivery on a "
        "first-seed book is still an em dash. Rank IC book did not move.",
    )
    add_h(doc, "PART 1 — How Roz came to be")
    add_h(doc, "1. Title and the question (60 seconds)")
    add_body(
        doc,
        "Title card from any Roz deck (Structured Narrative Research: "
        "Roz). One sentence: this is the summer close, not a weekly. "
        "The research question, from the 16 July deck — Motivation & "
        "Research Objective: does a disciplined, evidence-backed read "
        "of the earnings call add information beyond the point-in-time "
        "quant surprise?",
    )
    add_h(doc, "2. Three generations (90 seconds)")
    add_body(
        doc,
        "Pull: ECS Research Presentation slide 1 and Current Framework. "
        "Then Structured Narrative Research.pptx — The Objective / "
        "What each call is scored on. Then any Roz title.",
    )
    add_body(
        doc,
        "Earnings Call Summarizer was the origin. Structured Narrative "
        "is the scoring pipeline: dimensions, evidence ladder, "
        "narrative z, consolidated panel. Roz is the live surface "
        "on top of those artifacts — dashboard, Rank IC, Claims Desk, "
        "Post-Call Brief. Later layers did not replace earlier ones. "
        "Structured Narrative still produces what Roz reads.",
        ink=True,
    )
    add_h(doc, "3. The scoring contract (2 minutes)")
    add_body(
        doc,
        "Pull: 16 July — Evidence Verification, Point-in-Time & Feature "
        "Availability. 28 July — Evaluation Rigor / RankIC. Aug 19 — "
        "Instrumentation / Pre-registered Lab vs Production.",
    )
    add_body(
        doc,
        "Every claim is a transcript-backed excerpt. As of the August "
        "write-up: 26 companies, 950 scored quarters, 46 fiscal "
        "periods back to 2016-Q2, 14,855 excerpts, 100% "
        "transcript-supported. Ladder: verbatim 10,587 / composite "
        "3,357 / anchored 911 (plan-6eca5d31). Signals join a PIT "
        "quant spine. Rank IC is walk-forward. The unit of observation "
        "is a company-period, not pooled dimensions. Primary label is "
        "asof, not event-date fiscal labels. That was the July "
        "correction and it is still the contract.",
    )
    add_h(doc, "4. When the market adjusts (2.5 minutes)")
    add_body(
        doc,
        "Pull: 16 July agreement-effect story. Aug 19 Term Structure. "
        "Aug 27 Term Structure — Still Live. Do not say “novelty is "
        "U-shaped” without the era.",
    )
    add_body(
        doc,
        "Early read, 4-name pilot: change_magnitude is short-lived "
        "(+0.135 then flips). agrees_with_quant builds over the "
        "quarter. The market does not seem to price narrative/quant "
        "agreement on day one. That is why a live desk in the "
        "post-call window is the product, not a weekly summary.",
    )
    add_body(
        doc,
        "The print is not the adjustment. Agreement with the quant is "
        "absent in the first three weeks, so the work is the days after "
        "the call: open trees, this transcript, stronger or weaker, "
        "still silent. Rank IC still ranks the quarter. The live object "
        "is what you walk. Do not re-teach the two-curve term structure "
        "here — that chart already landed.",
        ink=True,
    )
    add_h(doc, "5. From a ranked quarter to a living object (3 minutes)")
    add_body(
        doc,
        "Pull: Aug 27 The New Object / The Bar — Autodesk Flex. "
        "Sep 3 The New Object — v2 / Claims Trees.",
    )
    add_body(
        doc,
        "Rank IC still ranks a quarter. It does not tell you whether "
        "last year’s promise is still open. Claims Desk is that object. "
        "v1: one row is type, clock, state, cite, follow-up. Kept is "
        "not delivered. Deliver rate is delivered / (delivered + "
        "missed). Unresolved is an em dash. Flex is the bar: ADSK "
        "FY2022-Q2 “we will launch Flex end of September,” FY2022-Q3 "
        "treats it as live usage mix. Kept and delivered. Path ID on "
        "that row is a miss. Path ID is not the verdict. Book deliver "
        "1 / 1 on the locked v1 book.",
    )
    add_body(
        doc,
        "v2 is a tree. One seed cite. Later quarters are edges — "
        "restated, silent, delivered, missed, expired. Walk may add "
        "silent. Walk does not invent delivered, hit, or missed. "
        "Three books: nvda_gold_v2, desk_ops_v2, desk_hc_v2. Gold: "
        "33 trees, 24 promises, 9 goals, window FY2022-Q2–FY2027-Q1, "
        "deliver 6 / 7 (H20 miss FY2026-Q3), hit 1 / 1, cue recall "
        "18 / 18. Gold file 84,215 bytes at the 27 August stamp.",
        ink=True,
    )
    add_h(doc, "6. Onboard, score, and keep walking (2 minutes)")
    add_body(
        doc,
        "Pull: 28 July Quartr Integration / Live Test FFIV. Aug 12 "
        "Live Loop / Host Automation. Aug 19 Universe Expansion.",
    )
    add_body(
        doc,
        "A new name is not “add a ticker to Rank IC.” Independent "
        "onboard is Quartr MCP only: search_companies → list_events → "
        "read_transcript, then --skip-pull. Earnings and conferences "
        "both seed. Overlay book for Independent is desk_hc_v2. After "
        "the call: walk the trees, Post-Call Brief, scorecard row. "
        "Then analyze: one company, the open trees, this call. Every "
        "terminal a person typed, with a cite. Unresolved stays a dash. "
        "Autopilot may insert a verbatim seed. It may not invent a close.",
    )
    add_h(doc, "7. Transparency, regimes, the dead window (2.5 minutes)")
    add_body(
        doc,
        "Pull: Sep 3 Claims Trees page. Live Roz: Management Regimes, "
        "Post-Call Brief. Do not call transparency desk_trust.",
    )
    add_body(
        doc,
        "Transparency (sidecar 2026-09-01T15:00:00+00:00) is a "
        "due-clock slip ledger. In-tree only. Someday wants with no "
        "clock are not ignore. Gold: 6 of 15 dated trees unanswered "
        "after the clock (40%). Expire closes and does not score. "
        "Silence is never a miss. Management Regimes: 74 CEO entries. "
        "Inherited vs closed-by-successor. Post-Call Brief is five "
        "questions on one screen in the five-day window before "
        "consensus moves: did they deliver, stronger or weaker, what "
        "resolved, what is still open or silent, who is management.",
        ink=True,
    )
    add_h(doc, "PART 2 — This week, and what landed yesterday")
    add_h(doc, "8. What the room already asked (60 seconds)")
    add_body(
        doc,
        "Pull: Sep 3 Immediate Next / The Room Picks. Do not re-teach "
        "term structure. One sentence: same 17 August book.",
    )
    add_body(
        doc,
        "Noon 31 August: tree not row, three books, gold 6 / 7, first "
        "seeds, leftover ~1,250, depth paused. Room bar was B — "
        "tighten the pull — then C. This week is what happened after "
        "that room.",
    )
    add_h(doc, "9. The afternoon of the 31st, then the pull (3 minutes)")
    add_body(
        doc,
        "Honest desk the same afternoon. Sector drives Book. Em dash "
        "is labeled “no scored X yet,” not 0%. Leftover cull 1,250 → "
        "1,145 (ops 556, HC 589). Gold recall stayed 1.0.",
    )
    add_body(
        doc,
        "Depth started, not closed. ISRG da Vinci X delivered "
        "FY2017-Q2 (11 of 166 systems were X; first clinical use in "
        "Germany; inside FY2017-Q4 clock). CRM $20B hit FY2021-Q4 "
        "(“over $20 billion in revenue”). LLY December 2016 dividend "
        "and MSFT BUILD still open.",
    )
    add_body(
        doc,
        "1 September: 43 tickers, 1,143 cue rows → 310 proposed seeds. "
        "12 / 310 not verbatim. Terminal pass on then-open trees: 21 "
        "actionable (7 delivered, 3 hit, 2 missed, 9 expired, 20 none). "
        "3 September workbook: 351 rows (41 verdicts + 310 seeds). "
        "Autopilot inserts what clears a hard bar. NVIDIA gold "
        "hands-off. After AVGO: ops 207 trees (155 confirmed, 52 "
        "provisional), healthcare 187 (141 confirmed, 46 provisional). "
        "That is the 44-skeleton book walked. Not a keep rate. "
        "Provisional is not gold.",
        ink=True,
    )
    add_h(doc, "10. The monitor was never watching (90 seconds)")
    add_body(
        doc,
        "2 September forensics: 8 events in five weeks, 7 "
        "manual_override=1. Watched set 5 of 45. Watch is now 45. "
        "Rank IC is still the 17 August comparison set. ADSK 27 August "
        "call never ran — that is the honest example.",
    )
    add_h(doc, "11. Case study — process, outcomes, meaning (6 minutes)")
    add_body(
        doc,
        "Paste CS-1 through CS-6 from documents/260910_case_study_slides.md. "
        "Demo last, not first. Leave-behind: Roz_Independent_Case_Study.html. "
        "CRWV stays on the desk as the younger-name precedent. Overlay "
        "book desk_hc_v2 is a filename, not a healthcare claim. Quartr "
        "MCP only. Delivery is an em dash, not 0%.",
        ink=True,
    )
    add_body(
        doc,
        "Process (CS-2). A new name is seven steps: pull FY + conferences "
        "through Quartr MCP; score eight dimensions on four stages "
        "(level, delta, surprise, novelty); join the PIT print z; write "
        "the brief; seed Claims Desk clocks; add a scorecard row; Rank IC "
        "stays parked. Conferences seed clocks. FY restates or ignores them.",
    )
    add_body(
        doc,
        "Outcomes (CS-3). DDOG: 28 FY + 37 conferences, 25/28 four-stage, "
        "200 panel rows, 65 scorecard rows, 10 overlay trees. LITE: 31 FY "
        "+ 26 conferences, 232 panel rows, 14 trees. First DDOG call "
        "12 Nov 2019. First LITE call 1 Nov 2018. Do not invent 2016. "
        "Both books: every open tree is silent. Delivery is a dash.",
        ink=True,
    )
    add_body(
        doc,
        "Meaning (CS-4 / CS-5 / CS-6). Datadog FY2026-Q2: demand +1.9 / "
        "surprise +1.2 / z −0.29 / gap 1.49 — the call is ahead of the "
        "print. Guidance +1.2 / delta −0.8 — Q3 28–29% from 36%. "
        "Confidence still +1.6; novelty −0.8 is the largest-customer "
        "caveat. Lumentum FY2026-Q4: demand agrees (+2.0 / z +0.15). "
        "Margins +2.0 / z −0.38 / gap 1.88. Cap alloc z +1.14 agrees. "
        "Macro −0.5 (Chinese InP). Live clock $2B / 40%. Dead clock "
        "$600M / 17–20%. Roz does not stamp bullish. Divergence is the "
        "use case. Slope vs adjective. Novelty is an object detector. "
        "A tree remembers what a quarter rank forgets.",
        ink=True,
    )
    add_body(
        doc,
        "This week’s chart, then the live click: Company Timeline plots "
        "every tracked call. Delivery stays a tooltip em dash. Claims "
        "Desk Book is All books. Company picker isolates DDOG or LITE. "
        "Then the HTML, then FY2026-Q2 / FY2026-Q4 Post-Call Brief.",
    )
    add_h(doc, "12. Still true / still not done (60 seconds)")
    add_body(
        doc,
        "Gold 6 / 7 and Flex 1 / 1 unchanged. LLY and MSFT BUILD "
        "open. Depth started, not closed. Healthcare is not a Rank IC. "
        "production_v1 frozen. Provisional ≠ gold. CRWV delivery is "
        "not 0%. Autopilot does not invent closes. 45/45 is not an "
        "unattended cycle.",
    )
    add_h(doc, "13. Thank you (45 seconds)")
    add_body(
        doc,
        "Pull any deck’s Thank You slide. Short. Guidance and advice "
        "all summer — the July Rank IC corrections, the asof contract, "
        "kept ≠ delivered, do not invent a close — that is the "
        "product. Then questions.",
    )
    add_h(doc, "Do not say")
    add_body(
        doc,
        "Ops, healthcare, or CRWV keep-rate / delivery 0%. Healthcare "
        "Rank IC. All twenty healthcare names scored gold. Provisional "
        "means gold. Kept means delivered. Horizon slip-miss means the "
        "tree missed. Path ID as the verdict. The remaining 125. "
        "Novelty is U-shaped (no era). 0_56 is the average of the "
        "buckets. Promotion, holdout, production_v1. Depth done. Flex "
        "undone. Unattended 45-name cycle. Invented delivery rates.",
    )
    add_h(doc, "Visual pull list")
    add_body(doc, "ECS Research Presentation.pptx — title, Current Framework.")
    add_body(
        doc,
        "Structured Narrative Research.pptx / 16 July — Motivation, "
        "Evidence Verification, PIT, Scale.",
    )
    add_body(
        doc,
        "Structured_Narrative_Research_Update___28_July_2026.pptx — "
        "RankIC, Composite, Quartr, Universe Expansion.",
    )
    add_body(
        doc,
        "Roz Research Presentaiton.pptx and Roz_Weekly_Update_2026-08-12.pptx "
        "— Roz dashboard, live loop, Narrative vs Quant.",
    )
    add_body(
        doc,
        "Roz August 19th.pptx — Workbench, Term Structure, What Survived / "
        "Died, Universe Expansion.",
    )
    add_body(
        doc,
        "Roz_Weekly_Update_2026-08-27.pptx — Term Structure, Claims Desk, "
        "Flex, Discipline.",
    )
    add_body(
        doc,
        "Roz_Weekly_Update_2026-09-03.pptx — v2 object, Claims Trees, "
        "Book control, The Room Picks.",
    )
    add_body(
        doc,
        "Live: localhost:8501/Claims_Desk — All books, Companies DDOG "
        "or LITE. Post_Call_Brief — DDOG FY2026-Q2, LITE FY2026-Q4 "
        "(default DDOG meeting is Citi TMT 8 Sep — switch period). "
        "Leave-behind: Roz_Independent_Case_Study.html. Optional gold "
        "6 / 7 on NVIDIA gold.",
    )
    add_h(doc, "Order")
    add_body(
        doc,
        "Title / question → three generations → scoring contract → "
        "when the market adjusts (the window) → Flex then trees → "
        "onboard, walk, analyze → transparency / dead window → [pause] → "
        "31 August recap → pull / 44 → hundreds → monitor → case study "
        "CS-1–CS-6 (names, process, book, Datadog gap, Lumentum split, "
        "meaning) → HTML / live desk → still true → thank you → questions. "
        "Demo at the end of Part 2, not at the open.",
    )
    dest = DESK / "Roz_Final_Presentation_Outline_2026-09-09.docx"
    doc.save(dest)
    return dest


def write_script() -> Path:
    doc = Document()
    style_doc(doc)
    add_title(doc, "Roz final script — 9 September 2026")
    add_body(
        doc,
        "Read this. Two parts. Do not invent a healthcare Rank IC. "
        "Do not call provisional trees gold. Do not give CRWV a "
        "delivery rate. Do not promote the 17 August book.",
    )

    add_h(doc, "PART 1 — Open")
    add_body(
        doc,
        "This is the last one. I want to do it in two parts. First, "
        "how Roz actually came to be — the process, not a feature "
        "list. Then what we built this week, including what landed "
        "tonight. I am not promoting a Rank IC book. Same asof "
        "stamp as 17 August, 17:28:40 UTC.",
    )
    add_body(
        doc,
        "The question has not changed since July. Does a disciplined, "
        "evidence-backed read of the earnings call add information "
        "beyond the point-in-time quant surprise? If it does, it "
        "has to show up before the street has already moved. That "
        "is the whole summer.",
    )

    add_h(doc, "Three generations")
    add_body(
        doc,
        "It started as the Earnings Call Summarizer. That was the "
        "scaffolding. Structured Narrative is the scoring engine on "
        "top of it: eight narrative dimensions, every claim tied to "
        "a transcript excerpt, joined to a point-in-time quant spine. "
        "Roz is the live layer. Dashboard, Rank IC, Claims Desk, "
        "Post-Call Brief. Roz does not replace Structured Narrative. "
        "It reads what that pipeline writes.",
    )

    add_h(doc, "The scoring contract")
    add_body(
        doc,
        "Two things the room taught me in July, and they are still "
        "the contract. First: the unit of observation is a company "
        "in a period. We do not pool dimensions to make a sample "
        "look bigger. Second: the primary cross-section is asof — "
        "the investable date — not a fiscal label that is not "
        "calendar-aligned. If I show you a Rank IC today, it is "
        "that construction.",
    )
    add_body(
        doc,
        "Every score is excerpt-backed. The August write-up: "
        "twenty-six companies, nine hundred fifty scored quarters, "
        "back to 2016 Q2, fourteen thousand eight hundred fifty-five "
        "excerpts, all transcript-supported. Verbatim, composite, "
        "or anchored. We do not keep unverified color.",
    )

    add_h(doc, "When the market adjusts")
    add_body(
        doc,
        "The first honest timing result was on the four-name pilot. "
        "A big tone change looks like short-term momentum that "
        "reverses. Agreement with the quant — narrative and the "
        "print pointing the same way — does not show up in the "
        "first three weeks. It builds. That is the opposite of a "
        "headline that is already in the price. That is why the "
        "desk has to be live in the days after the call, not a "
        "memo the next month.",
    )
    add_body(
        doc,
        "I am not re-teaching the two-curve term structure here. "
        "The print is not the adjustment. Agreement builds over "
        "the quarter. The useful work is the days after the call: "
        "open trees, this transcript, what got stronger or weaker, "
        "what is still silent. Rank IC still ranks the quarter. "
        "The live object is what you walk.",
    )

    add_h(doc, "From a ranked quarter to a living object")
    add_body(
        doc,
        "Rank IC tells you how a quarter ranks. It does not tell "
        "you whether last year’s promise is still open. That is "
        "why Claims Desk exists. A kept object is the same claim "
        "coming back next quarter. Delivered means a dated promise "
        "came true. Those are different words. Unresolved is a "
        "dash. It is not zero percent.",
    )
    add_body(
        doc,
        "The bar is still Autodesk Flex. FY2022 Q2: we will launch "
        "Flex at the end of September. FY2022 Q3 they talk about "
        "it as live usage mix — net new, occasional, advanced. "
        "Kept and delivered. Path ID on that row is a miss. Path "
        "ID is not the verdict. Locked v1 book deliver is one over "
        "one, that tree only.",
    )
    add_quote(
        doc,
        "At the end of September, we will launch a new pay-as-you-go "
        "consumption model, called Flex. It matches the customer's "
        "cost with their usage.",
    )
    add_body(
        doc,
        "v2 is a tree, not a row. One seed cite. Later quarters are "
        "edges. The walker may mark silent. It may not invent "
        "delivered, hit, or missed. NVIDIA gold is locked. "
        "Thirty-three trees. Cue recall one hundred percent. "
        "Deliver six over seven. The miss is H20 in FY2026 Q3. "
        "Tech ops and healthcare are operational books. First "
        "seeds used to read as a dash, not a keep rate. That was "
        "the honest status on 31 August.",
    )

    add_h(doc, "Onboard, score, keep walking")
    add_body(
        doc,
        "A new company is not “drop it into Rank IC.” Independent "
        "onboard pulls transcripts through Quartr in Cursor — "
        "search, list events, read the transcript — then we skip "
        "a second pull. Earnings and conferences both count. The "
        "overlay book for an Independent is the healthcare claims "
        "book because that is the overlay file, not because the "
        "name is healthcare. After each call we walk, write a "
        "brief, and add a scorecard row. Then we analyze: one "
        "company, the open trees, this call. Unresolved stays a "
        "dash. The machine may propose a seed. It does not close "
        "a tree from silence.",
    )

    add_h(doc, "Transparency and the dead window")
    add_body(
        doc,
        "Transparency is a narrower question. After a dated clock, "
        "did they mention the claim again? Gold: six of fifteen "
        "dated trees had no later cite. Forty percent. That is not "
        "desk trust. Expire closes a clock without scoring it. "
        "Silence is not a miss. There is a Management Regimes page "
        "so a promise seeded under one CEO can be inherited, "
        "adopted, or ignored by the next. Seventy-four regimes.",
    )
    add_body(
        doc,
        "The product in the room is the five-day window. Consensus "
        "has not revised yet. Post-Call Brief is five questions on "
        "one screen: did they deliver, are forward commitments "
        "stronger or weaker, what resolved, what is still open or "
        "silent, and who is on the hook. Claims Desk holds the "
        "trees. Rank IC still ranks the quarter. They are different "
        "objects. That is Part 1.",
    )

    add_h(doc, "PART 2 — What the room asked")
    add_body(
        doc,
        "Same Rank IC book. I am not re-teaching term structure. "
        "Last time we met, around noon on the 31st, I showed you a "
        "tree, three books, and NVIDIA gold at six over seven. "
        "Tech and healthcare were first seeds. The leftover dump "
        "was the hole. You said tighten the pull, then go deep on "
        "four names. This week is what happened after that room.",
    )

    add_h(doc, "The afternoon, then the machine")
    add_body(
        doc,
        "Honest desk rode that afternoon. Sector now picks the book. "
        "An em dash is labeled as no scored closes yet, not zero. "
        "The leftover list went from twelve hundred fifty openings "
        "to eleven hundred forty-five. That is a shorter worklist. "
        "It is not a short one. Gold recall stayed one.",
    )
    add_body(
        doc,
        "We un-paused the four-name file. Intuitive da Vinci X is "
        "delivered in FY2017 Q2. Eleven of one hundred sixty-six "
        "systems were X, first clinical use in Germany, inside the "
        "clock. Salesforce over twenty billion is a hit in FY2021 "
        "Q4, not the COVID-year guide. Lilly’s December 2016 "
        "dividend is still open. Microsoft BUILD is still open. I "
        "will not close those from silence.",
    )
    add_body(
        doc,
        "Then the pull became a machine. On the first we ran one "
        "model pass per name over the leftover cues. Forty-three "
        "tickers. Three hundred ten proposed seeds. Twelve of those "
        "quotes were not verbatim. They do not get in. Twenty-one "
        "actionable closes on the then-open trees. Propose only.",
    )
    add_body(
        doc,
        "On the third we reviewed all of them in one workbook. "
        "Three hundred fifty-one rows. High-confidence seeds were "
        "accepted. Mediums went provisional. Then the autopilot "
        "was allowed to insert what clears a hard bar: verbatim, "
        "valid tree, not a duplicate sentence. One dollar per "
        "ticker. Gold is off limits. The book you saw on the 31st "
        "had forty-four skeleton trees. After AVGO: two hundred "
        "seven ops trees, one hundred eighty-seven healthcare. "
        "One hundred fifty-five and one hundred forty-one of those "
        "are confirmed. The rest are provisional. That is not a "
        "keep rate. That is a walked desk.",
    )

    add_h(doc, "The monitor")
    add_body(
        doc,
        "I owe you an honest sentence on automation. In five weeks "
        "the live monitor database had eight events. Seven were "
        "typed in by hand. We were watching five companies out of "
        "forty-five. Watch is now forty-five. Rank IC is still the "
        "17 August book. Autodesk’s 27 August call never ran. That "
        "is the example. This is not an unattended forty-five-name "
        "cycle.",
    )

    add_h(doc, "Case study — process")
    add_body(
        doc,
        "The live names are Datadog and Lumentum. Independent. "
        "Industry tag is tech. They sit on the healthcare overlay "
        "because that is the file, not because they are drug "
        "companies. CoreWeave stays on the desk. It is not the hero. "
        "They are not in the 17 August Rank IC book.",
    )
    add_body(
        doc,
        "Here is what actually runs on a new name. We pull every "
        "earnings call and every conference through Quartr in "
        "Cursor — Q&A tails included. Structured Narrative scores "
        "eight dimensions on four stages: how this quarter looks, "
        "whether the story got better or worse versus last call, "
        "whether they were more bullish than the Street, and "
        "whether they said something new. That joins a point-in-time "
        "print z — this company’s own surprise history, as of the "
        "call date. Then a brief, then Claims Desk clocks, then a "
        "scorecard row for this meeting. Rank IC stays parked. "
        "Conferences are where a lot of clocks are first said. "
        "FY is where they get restated or ignored.",
    )

    add_h(doc, "Case study — outcomes")
    add_body(
        doc,
        "Datadog: twenty-eight earnings transcripts, thirty-seven "
        "conferences. First call 12 November 2019. Twenty-five of "
        "twenty-eight FY quarters have all four stages. Two hundred "
        "panel rows. Sixty-five scorecard rows. Ten overlay trees. "
        "Latest quarter FY2026 Q2, sixth of August. Lumentum: "
        "thirty-one earnings, twenty-six conferences. First call "
        "1 November 2018. June-thirtieth fiscal. Two hundred "
        "thirty-two panel rows. Fourteen trees. Latest quarter "
        "FY2026 Q4, eleventh of August. Do not invent 2016. Quartr "
        "has nothing earlier. Delivery on both books is a dash. "
        "Ten silent and fourteen silent. Do not say zero percent. "
        "Tree count is inventory, not a keep rate.",
    )

    add_h(doc, "Case study — meaning")
    add_body(
        doc,
        "Datadog just ran a confident call on a slightly negative "
        "print z. Demand is plus 1.9, surprise plus 1.2, z minus "
        "0.29, gap 1.49. They sold thirty-six percent growth and "
        "record sequential adds. The call is ahead of the number. "
        "Guidance is still plus 1.2 on the adjective — they raised "
        "dollars — and minus 0.8 on the slope. Q3 growth twenty-eight "
        "to twenty-nine from thirty-six, after they de-risked the "
        "largest customer. Confidence did not collapse. Novelty "
        "minus 0.8 is the new caveat. FedRAMP High is still open. "
        "Bits SRE went GA. India/Brazil is due this quarter. Ask it.",
    )
    add_body(
        doc,
        "Lumentum is the contrast. A billion-oh-one quarter, one "
        "twenty-five guided. Demand agrees with the print. Cash "
        "agrees — they retired 1.1 billion of converts, z plus 1.14. "
        "Margins do not. Level pinned at plus 2.0, surprise plus 1.5, "
        "z still minus 0.38, gap 1.88. They crossed fifty percent "
        "gross a year early versus their own two-billion model. "
        "Macro is the only red cell, minus 0.5. Chinese indium "
        "phosphide just entered the script. The six-hundred-million "
        "/ seventeen-to-twenty clock is already dead. The live clock "
        "is two billion / forty percent in eighteen to twenty-four "
        "months.",
    )
    add_body(
        doc,
        "Roz does not stamp bullish. Divergence is the use case — "
        "the call saying something the print z does not. Slope "
        "versus adjective. Novelty is an object detector, not a "
        "louder tone. A tree remembers what a quarter rank forgets. "
        "On XLK, not these names, raw tone reversed. Agreement was "
        "flat three weeks and built later. These two show the "
        "machinery on a live print. They are not a live IC.",
    )
    add_body(
        doc,
        "This week the scorecard chart stopped throwing away any "
        "call without a delivery rate. An Independent used to "
        "vanish. It now plots every tracked call, in time. "
        "Delivery is a dash in the tooltip. We do not draw that "
        "dash as zero. Tonight Claims Desk defaults to All books. "
        "The company picker isolates DDOG or LITE. The leave-behind "
        "is the HTML. Then Independent, DDOG or LITE, Claims Desk, "
        "Company Timeline, FY Post-Call Brief. Switch DDOG off the "
        "Citi TMT conference onto FY2026 Q2.",
    )

    add_h(doc, "Still true")
    add_body(
        doc,
        "Gold is still six over seven. Flex is still one over one. "
        "Lilly and Microsoft BUILD are still open. Depth is started, "
        "not closed. Healthcare is not a Rank IC. I have not "
        "promoted production_v1. Provisional trees are not gold. "
        "CoreWeave is not zero percent delivered.",
    )

    add_h(doc, "Thank you")
    add_body(
        doc,
        "I want to say thank you. The useful parts of this are "
        "the things you pushed on all summer. Unit of observation. "
        "Asof, not a messy fiscal label. Kept is not delivered. "
        "Do not invent a close. Do not call a dash zero. That "
        "guidance is the product as much as the software. I "
        "appreciate the time and the advice. I am happy to take "
        "questions. If you want the live object: the HTML first, "
        "then Independent, DDOG or LITE, Claims Desk, Company "
        "Timeline, then the FY Post-Call Brief.",
    )

    dest = DESK / "Roz_Final_Presentation_Script_2026-09-09.docx"
    doc.save(dest)
    return dest


def main() -> int:
    DESK.mkdir(parents=True, exist_ok=True)
    for path in (write_outline(), write_script()):
        print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
